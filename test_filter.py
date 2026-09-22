"""Cross-check /filter against independently written SQL.

The point is not that filter_rows() runs, but that it returns the SAME set a
differently-phrased query returns. Every expectation below is computed from
scratch rather than reusing app.py's own SQL, so a bug in filter_rows cannot
hide by being consistent with itself.
"""
import os
import shutil
import sqlite3
import tempfile

import app as A

A.DB = os.path.join(tempfile.mkdtemp(), "t.db")
shutil.copy(os.path.join(os.path.dirname(os.path.abspath(__file__)), "ukp.db"), A.DB)
A.LOG_DB = A.DB
A.app.config["TESTING"] = True
A.app.secret_key = "test"

con = sqlite3.connect(A.DB)
con.row_factory = sqlite3.Row
BIG = 10 ** 9
ctx = A.app.test_request_context()
ctx.push()

# ---------------------------------------------------------------- helpers
def rows(**kw):
    kw.setdefault("limit", BIG)
    return A.filter_rows(**kw)


def keys(rs):
    """Identity of a result row = the registration it stands for."""
    return {(r["sc"], r["ijzh"], r["last_ex"]) for r in rs}


def sql1(q, a=()):
    return con.execute(q, a).fetchone()[0]


# ---------------------------------------------------------------- 1. no filter
all_rows = rows()
expect = sql1("SELECT COUNT(*) FROM (SELECT p.uc FROM peserta p"
              " JOIN nilai n ON n.uc=p.uc GROUP BY p.uc)")
assert len(all_rows) == expect, f"unfiltered {len(all_rows)} != {expect}"
print(f"1. unfiltered          : {len(all_rows)} registrations")

# no row may be a duplicate registration
assert len({(r["sc"], r["ijzh"], r["last_ex"], r["att"]) for r in all_rows}) == len(all_rows), \
    "duplicate rows in output"

# ---------------------------------------------------------------- 2. lulus split
lulus = rows(status="lulus")
belum = rows(status="belum_lulus")
assert len(lulus) + len(belum) == len(all_rows), \
    f"{len(lulus)}+{len(belum)} != {len(all_rows)} - statuses must partition"
assert not (keys(lulus) & keys(belum)), "a registration is both lulus and belum"
assert all(r["lulus"] for r in lulus) and not any(r["lulus"] for r in belum), \
    "lulus flag disagrees with the status filter"
exp_l = sql1("SELECT COUNT(*) FROM (SELECT p.uc FROM peserta p JOIN nilai n ON n.uc=p.uc"
             " GROUP BY p.uc HAVING SUM(n.lulus)>0)")
assert len(lulus) == exp_l, f"lulus {len(lulus)} != {exp_l}"
print(f"2. lulus/belum split   : {len(lulus)} + {len(belum)} = {len(all_rows)}")

# ---------------------------------------------------------------- 3. belum_mengulang
bm = rows(status="belum_mengulang")
assert keys(bm) <= keys(belum), "belum_mengulang must be a subset of belum_lulus"

# independent re-derivation: for each failed registration, is there a later
# exam at the same level by the same person?
exp_bm = set()
for r in con.execute("""
        SELECT p.sc, p.ijzh, MAX(n.tgl_ujian) last_ex
        FROM peserta p JOIN nilai n ON n.uc=p.uc
        GROUP BY p.uc HAVING SUM(n.lulus)=0"""):
    later = sql1("""SELECT COUNT(*) FROM peserta p2 JOIN nilai n2 ON n2.uc=p2.uc
                    WHERE p2.sc=? AND p2.ijzh=? AND n2.tgl_ujian>?""",
                 (r["sc"], r["ijzh"], r["last_ex"]))
    if not later:
        exp_bm.add((r["sc"], r["ijzh"], r["last_ex"]))
assert keys(bm) == exp_bm, (f"belum_mengulang mismatch: got {len(bm)}, "
                            f"expected {len(exp_bm)}, "
                            f"diff {list(keys(bm) ^ exp_bm)[:3]}")
print(f"3. belum_mengulang     : {len(bm)} (independently re-derived, exact match)")

# this is the number the standalone report produced - they must agree
print(f"   report_no_retake.py agreement: {len(bm)} == 677 -> {len(bm) == 677}")

# ---------------------------------------------------------------- 4. belum_skl
bs = rows(status="belum_skl")
assert all(r["lulus"] for r in bs), "belum_skl contains a non-passing registration"
assert all(r["skl_status"] == "BELUM TERBIT" for r in bs), \
    "belum_skl row does not read BELUM TERBIT"
exp_bs = sql1("""SELECT COUNT(*) FROM (SELECT p.uc FROM peserta p
    JOIN nilai n ON n.uc=p.uc WHERE p.uc NOT IN (SELECT uc FROM skl)
    GROUP BY p.uc HAVING SUM(n.lulus)>0)""")
assert len(bs) == exp_bs, f"belum_skl {len(bs)} != {exp_bs}"
print(f"4. belum_skl           : {len(bs)}")

# ---------------------------------------------------------------- 5. diklat
for d in ("AMC", "STIP", "AMD"):
    got = rows(diklat=d)
    exp = sql1("SELECT COUNT(*) FROM (SELECT p.uc FROM peserta p"
               " JOIN nilai n ON n.uc=p.uc WHERE p.diklat=? GROUP BY p.uc)", (d,))
    assert len(got) == exp, f"diklat {d}: {len(got)} != {exp}"
    assert all(r["diklat"] == d for r in got), f"diklat {d} leaked another institution"
print("5. diklat filter       : AMC/STIP/AMD all match")

# the session's original question, end to end
amc = rows(diklat="AMC", status="belum_mengulang")
assert len(amc) == 36, f"AMC belum_mengulang = {len(amc)}, expected 36"
assert all(r["diklat"] == "AMC" and not r["lulus"] for r in amc)
print(f"   AMC + belum_mengulang: {len(amc)} (matches this session's report)")

# ---------------------------------------------------------------- 6. ijzh + tahun
for lv in ("PASN3", "GMDSS", "UGN5"):
    got = rows(ijzh=lv)
    exp = sql1("SELECT COUNT(*) FROM (SELECT p.uc FROM peserta p"
               " JOIN nilai n ON n.uc=p.uc WHERE p.ijzh=? GROUP BY p.uc)", (lv,))
    assert len(got) == exp, f"ijzh {lv}: {len(got)} != {exp}"
    assert all(r["ijzh"] == lv for r in got), f"ijzh {lv} leaked another level"
print("6. ijzh filter         : PASN3/GMDSS/UGN5 all match")

for y in ("2026", "2023", "2021"):
    got = rows(tahun=y)
    assert all(r["last_ex"][:4] == y for r in got), f"tahun {y} leaked another year"
    exp = sql1("SELECT COUNT(*) FROM (SELECT p.uc FROM peserta p JOIN nilai n ON n.uc=p.uc"
               " GROUP BY p.uc HAVING substr(MAX(n.tgl_ujian),1,4)=?)", (y,))
    assert len(got) == exp, f"tahun {y}: {len(got)} != {exp}"
print("7. tahun filter        : 2026/2023/2021 all match (on LAST exam, not any)")

# years must partition the whole set
ysum = sum(len(rows(tahun=y)) for y in
           [r[0] for r in con.execute(
               "SELECT DISTINCT substr(tgl_ujian,1,4) FROM nilai"
               " WHERE tgl_ujian IS NOT NULL")])
assert ysum == len(all_rows), f"years sum to {ysum}, not {len(all_rows)}"
print(f"   years partition      : {ysum} == {len(all_rows)}")

# ---------------------------------------------------------------- 8. combined
combo = rows(diklat="AMC", ijzh="PASN3", status="belum_lulus")
manual = [r for r in all_rows
          if r["diklat"] == "AMC" and r["ijzh"] == "PASN3" and not r["lulus"]]
assert keys(combo) == keys(manual), \
    f"combined filter {len(combo)} != manual intersection {len(manual)}"
print(f"8. combined filter     : AMC+PASN3+belum_lulus = {len(combo)} "
      f"(equals manual intersection)")

# order of filters must not matter
assert keys(rows(ijzh="PASN3", diklat="AMC", status="belum_lulus")) == keys(combo), \
    "filter order changes the result"

# ---------------------------------------------------------------- 9. paging
lim = A.filter_rows(status="belum_mengulang", limit=10)
assert len(lim) == 10, f"limit ignored: got {len(lim)}"
assert keys(lim) <= keys(bm), "limited page contains rows outside the full set"
assert A.filter_count(status="belum_mengulang") == len(bm), "filter_count disagrees"

# Walking every page must reproduce the full set exactly - no row seen twice,
# none missed. This is the failure paging actually has, so test it on real data.
per, seen, order = 100, [], []
for pg in range(0, 10):
    chunk = A.filter_rows(status="belum_mengulang", limit=per, offset=pg * per)
    if not chunk:
        break
    assert len(chunk) <= per, f"page {pg} returned {len(chunk)} rows, max {per}"
    seen.extend(keys(chunk))
    order.extend(r["last_ex"] for r in chunk)
assert len(seen) == len(bm), f"paging yielded {len(seen)} rows, full set is {len(bm)}"
assert len(set(seen)) == len(seen), "a row appears on more than one page"
assert set(seen) == keys(bm), "paged set differs from the unpaged set"
assert order == sorted(order, reverse=True), "sort order breaks across page boundaries"
print(f"9. paging              : {len(bm)} rows over {-(-len(bm) // per)} pages, "
      f"no duplicates, no gaps, order intact")

# Repeating the same page must return the same rows - the uc tiebreak matters
# because many registrations share a last_ex date.
a1 = A.filter_rows(status="belum_mengulang", limit=per, offset=per)
a2 = A.filter_rows(status="belum_mengulang", limit=per, offset=per)
assert keys(a1) == keys(a2), "same page returns different rows on repeat"

# and on a big unfiltered set too, where ties are far more common
big = []
for pg in range(0, 6):
    big.extend(keys(A.filter_rows(limit=200, offset=pg * 200)))
assert len(set(big)) == len(big), "duplicate rows across pages on the unfiltered set"
print(f"   tie stability       : repeated page identical; "
      f"{len(big)} unfiltered rows unique across 6 pages")

# offset past the end is empty, not an error
assert A.filter_rows(status="belum_mengulang", limit=per, offset=10 ** 6) == []

# count must not depend on paging
assert A.filter_count(diklat="AMC", status="belum_mengulang") == 36

# empty result must not explode
assert rows(diklat="AMC", ijzh="UGN1") == [], "impossible combination returned rows"
print("   empty combination    : returns [] cleanly")

# bad status falls back rather than erroring
assert len(rows(status="not-a-status")) == len(all_rows), "unknown status not ignored"

ctx.pop()

# ---------------------------------------------------------------- 10. HTTP layer
# The test mints its own admin credential rather than reading the real one, so
# it never depends on (or reveals) the production password.
import re
from werkzeug.security import generate_password_hash

TEST_PW = "filter-test-only-pw"
A.ADMIN_HASH = generate_password_hash(TEST_PW)

with A.app.test_client() as c:
    assert c.get("/filter").status_code == 302, "filter not gated for guests"
    p = c.get("/login").get_data(as_text=True)
    tok = re.search(r'name="csrf" value="([^"]+)"', p).group(1)
    r = c.post("/login", data={"user": A.ADMIN_USER, "pw": TEST_PW, "csrf": tok})
    assert r.status_code == 302, f"login failed: {r.status_code}"

    page = c.get("/filter").get_data(as_text=True)
    for d in ("AMC", "STIP"):
        assert f'value="{d}"' in page, f"diklat {d} missing from dropdown"
    assert "belum_mengulang" in page, "status option missing"

    tok = re.search(r'name="csrf" value="([^"]+)"', page).group(1)
    h = c.post("/filter", data={"diklat": "AMC", "status": "belum_mengulang",
                                "csrf": tok}).get_data(as_text=True)
    assert "dari <b>36</b> registrasi" in h, "page does not report 36 AMC rows"
    for sc in ("6211928464", "6211406146"):
        assert sc in h, f"known AMC dormant code {sc} missing from page"
    assert "BELUM LULUS" in h and "LULUS</span>" in h
    # 36 < 100, so a single page and no pager
    assert 'class="pager"' not in h, "pager shown for a single-page result"

    # A multi-page result: walk it over HTTP and confirm the browser sees
    # every row exactly once, with continuous numbering.
    n = A.filter_count(status="belum_mengulang")
    pages = -(-n // A.PER_PAGE)
    codes, nums = [], []
    for p in range(1, pages + 1):
        body = c.get("/filter", query_string={"status": "belum_mengulang",
                                              "page": p}).get_data(as_text=True)
        rowhtml = body.split("<tbody>")[1].split("</tbody>")[0]
        cells = re.findall(r"<tr>\s*<td>(\d+)</td>\s*<td><a[^>]*>(\d+)</a>", rowhtml)
        assert cells, f"page {p} rendered no rows"
        nums.extend(int(x[0]) for x in cells)
        codes.extend(x[1] for x in cells)
    assert nums == list(range(1, n + 1)), \
        f"row numbering not continuous 1..{n} across {pages} pages"
    assert len(codes) == n and len(set(zip(codes, nums))) == n, \
        "HTTP paging dropped or repeated a row"
    print(f"    paging over HTTP   : {n} rows, {pages} pages, numbering 1..{n} continuous")

    # page beyond the end clamps instead of erroring or showing nothing
    last = c.get("/filter", query_string={"status": "belum_mengulang",
                                          "page": 9999}).get_data(as_text=True)
    assert f"halaman {pages} dari {pages}" in last, "out-of-range page did not clamp"
    bad = c.get("/filter", query_string={"status": "belum_mengulang",
                                         "page": "abc"})
    assert bad.status_code == 200 and "halaman 1 dari" in bad.get_data(as_text=True), \
        "non-numeric page crashed the view"
    print("    page clamping      : 9999 -> last page, 'abc' -> page 1")

    # the pager must carry the filter, not drop it
    p2 = c.get("/filter", query_string={"diklat": "STIP", "status": "belum_lulus",
                                        "page": 2}).get_data(as_text=True)
    assert "diklat=STIP" in p2 and "status=belum_lulus" in p2, \
        "pager links lose the active filter"

    x = c.post("/filter.xlsx", data={"diklat": "AMC", "status": "belum_mengulang",
                                     "csrf": tok})
    assert x.status_code == 200 and len(x.data) > 3000, "xlsx export failed"
    assert x.headers["Content-Disposition"].startswith("attachment")
    assert "belum_mengulang-AMC" in x.headers["Content-Disposition"]
    print(f"10. HTTP + export      : page OK, xlsx {len(x.data)} bytes")

    # the export must contain exactly the rows the filter selected
    import io
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(x.data)).active
    hdr = next(i for i, r in enumerate(ws.iter_rows(values_only=True), 1)
               if r and r[0] == "No")
    body = [r for r in list(ws.iter_rows(values_only=True))[hdr:]
            if r and isinstance(r[0], int)]
    assert len(body) == 36, f"xlsx has {len(body)} data rows, expected 36"
    assert {r[1] for r in body} == {r["sc"] for r in amc}, \
        "xlsx seafarer codes differ from the filtered set"
    assert all(r[9] == "BELUM LULUS" for r in body), "xlsx status column wrong"
    print(f"    xlsx contents      : {len(body)} rows, codes match the filter exactly")

    c.get("/logout")
    assert c.get("/filter").status_code == 302, "logout did not drop access"

print("\nfilter recheck: all cross-checks passed")
con.close()
