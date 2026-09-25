"""Extract the competency workbook into kompetensi.json.

Run once when 'Daftar Kompetensi UKP Perdana.xlsx' is updated:
    .venv\\Scripts\\python.exe build_kompetensi.py

The app reads the JSON, never the workbook: openpyxl on every request would
be slow, and the file lives outside the deployed directory.
"""
import json
import os
import sys

import openpyxl

SRC = os.environ.get(
    "KOMPETENSI_XLSX",
    r"D:\PUKP-3\Daftar Kompetensi UKP Perdana.xlsx")
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "kompetensi.json")

DETAIL = ["Nautika Peningkatan", "Nautika Pembentukan", "Keterampilan",
          "Teknika Peningkatan", "Teknika Pembentukan"]

# Workbook level name -> the ijzh code family used in ukp.db / the dashboard.
CODE = {
    "ANT I": "UGN1", "ANT II": "UGN2", "ANT III": "UGN3",
    "ANT IV": "UGN4", "ANT V": "UGN5",
    "ANT III Pra Prala": "PRAN3", "ANT III Pasca Prala": "PASN3",
    "ANT II Pra Layar": "PRAN2", "ANT II Pasca Layar (D-IV Lanjutan)": "PASN2",
    "GMDSS": "GMDSS",
    "ATT I": "UGT1", "ATT II": "UGT2", "ATT III": "UGT3",
    "ATT IV": "UGT4", "ATT V": "UGT5",
    "ATT III Pra Prala": "PRAT3", "ATT III Pasca Prala": "PAST3",
    "ATT II Pra Layar": "PRAT2", "ATT II Pasca Layar (D-IV Lanjutan)": "PAST2",
}


def txt(c):
    return "" if c is None else str(c).strip()


def main():
    if not os.path.exists(SRC):
        sys.exit(f"tidak ditemukan: {SRC}")
    wb = openpyxl.load_workbook(SRC, data_only=True)

    # ---- index sheet: level -> full certificate name + declared subject count
    idx, group = {}, None
    for r in wb["Daftar Kompetensi UKP"].iter_rows(min_row=4, values_only=True):
        a, b, c, d = (txt(x) for x in (r + (None,) * 4)[:4])
        if a and not b and not c:
            if a.upper().startswith("TOTAL"):
                continue
            group = a.split(". ", 1)[-1]
            continue
        if b and c:
            idx[b] = {"nama": c, "n": int(d or 0), "grup": group}

    # ---- recap sheet: CBA / Komprehensif split
    rec = {}
    for r in wb["Rekap Mata Uji"].iter_rows(min_row=4, values_only=True):
        a, b, c, d, e = (txt(x) for x in (r + (None,) * 5)[:5])
        if b and c != "":
            rec[b] = {"cba": int(c or 0), "komp": int(d or 0), "total": int(e or 0)}

    # ---- detail sheets: the subjects themselves
    levels, order = {}, []
    for sh in DETAIL:
        ws = wb[sh]
        grup = txt(ws["A1"].value).split(" - ", 1)[-1].split(". ", 1)[-1]
        cur = None
        for r in ws.iter_rows(min_row=3, values_only=True):
            a, b, c, d, e, f = (txt(x) for x in (r + (None,) * 6)[:6])
            if a and not b and not c:
                if a.lower().startswith("subtotal"):
                    cur = None
                    continue
                cur = a.split(" — ", 1)[0].strip()      # "ANT I — AHLI ..." -> "ANT I"
                if cur not in levels:
                    order.append(cur)
                    levels[cur] = {"tingkat": cur, "grup": grup, "mu": []}
                continue
            if cur and c and e:
                levels[cur]["mu"].append(
                    {"kode": c, "no": int(d or 0), "nama": e, "jenis": f or "CBA"})

    # ---- stitch + verify against the workbook's own totals
    out, bad = [], []
    for lv in order:
        d = levels[lv]
        meta, rc = idx.get(lv, {}), rec.get(lv, {})
        mu = d["mu"]
        # The recap sheet has only two columns, so it files GMDSS's two
        # 'Praktik' rows under Komprehensif. Keep the workbook's own label on
        # each subject, but count anything that is not CBA as non-CBA so the
        # totals still reconcile with 'Rekap Mata Uji'.
        cba = sum(1 for m in mu if m["jenis"].upper().startswith("CBA"))
        komp = len(mu) - cba
        if meta.get("n") and meta["n"] != len(mu):
            bad.append(f"{lv}: indeks {meta['n']} != {len(mu)} baris")
        if rc:
            if rc["total"] != len(mu):
                bad.append(f"{lv}: rekap total {rc['total']} != {len(mu)}")
            if rc["cba"] != cba:
                bad.append(f"{lv}: rekap CBA {rc['cba']} != {cba}")
            if rc["komp"] != komp:
                bad.append(f"{lv}: rekap komprehensif {rc['komp']} != {komp}")
        nos = [m["no"] for m in mu]
        if nos != list(range(1, len(nos) + 1)):
            bad.append(f"{lv}: nomor mata uji tidak berurutan")
        out.append({
            "tingkat": lv, "kode": CODE.get(lv, ""), "grup": d["grup"],
            "nama": meta.get("nama", lv), "cba": cba, "komp": komp,
            "total": len(mu), "mu": mu,
        })

    if bad:
        print("SELISIH dengan rekap workbook:")
        for b in bad:
            print("  -", b)
        sys.exit(1)

    tot = sum(x["total"] for x in out)
    doc = {"sumber": os.path.basename(SRC), "tingkat": out,
           "total_mu": tot, "total_tingkat": len(out),
           "total_cba": sum(x["cba"] for x in out),
           "total_komp": sum(x["komp"] for x in out)}
    with open(OUT, "w", encoding="utf-8") as fh:
        json.dump(doc, fh, ensure_ascii=False, indent=1)

    print(f"{len(out)} tingkat ijazah, {tot} mata uji -> {OUT}")
    for x in out:
        miss = "" if x["kode"] else "  <- TIDAK ADA KODE"
        print(f"  {x['tingkat']:36s} {x['kode']:6s} {x['total']:3d} mu"
              f"  ({x['cba']} CBA + {x['komp']} komp){miss}")
    print("cocok dengan sheet 'Rekap Mata Uji' dan 'Daftar Kompetensi UKP'")


if __name__ == "__main__":
    main()
