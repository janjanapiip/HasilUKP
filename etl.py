"""Load ukp_v2021.xlsx + ukp_v2026.xlsx into a single SQLite db.

Run:  python etl.py [--out ukp.db]
Idempotent: drops and rebuilds the data tables every run.
Writes exceptions.csv listing rows that could not be attached to a person.
"""
import argparse
import csv
import datetime
import json
import os
import re
import sqlite3
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
SRC_DIR = os.path.join(os.path.dirname(HERE), "v2026")
SOURCES = [("2021", "ukp_v2021.xlsx"), ("2026", "ukp_v2026.xlsx")]  # 2026 wins on conflict

# Birth dates that appear hundreds of times: data-entry placeholders, not real.
PLACEHOLDER_DOB = {
    datetime.date(1996, 6, 5),
    datetime.date(2003, 1, 15),
}

SCHEMA = """
DROP TABLE IF EXISTS peserta;
DROP TABLE IF EXISTS nilai;
DROP TABLE IF EXISTS skl;
DROP TABLE IF EXISTS ref_ijzh;
DROP TABLE IF EXISTS ref_diklat;

CREATE TABLE peserta (
    uc         TEXT PRIMARY KEY,
    sc         TEXT NOT NULL,
    noukp      TEXT,
    nama       TEXT,
    tpt_lahir  TEXT,
    tgl_lahir  TEXT,          -- ISO date or NULL when unusable/placeholder
    dob_usable INTEGER NOT NULL,
    diklat     TEXT,
    ijzh       TEXT,
    ukp1       TEXT,          -- ISO date or NULL
    src        TEXT NOT NULL
);
CREATE INDEX ix_peserta_sc ON peserta(sc);
CREATE INDEX ix_peserta_nama ON peserta(nama);

CREATE TABLE nilai (
    id           INTEGER PRIMARY KEY,
    uc           TEXT NOT NULL,
    uc_raw       TEXT NOT NULL,  -- key as written in the sheet, before matching
    ijzh         TEXT,
    nama         TEXT,
    tgl_ujian    TEXT,
    tgl_sidang   TEXT,
    mengulang_ke INTEGER,
    her          INTEGER,     -- source's count of MU below 70
    her_calc     INTEGER,     -- recomputed; mismatch means dirty source row
    lulus        INTEGER,
    mu           TEXT NOT NULL,  -- JSON list, length = active MU count for this ijzh
    src          TEXT NOT NULL,
    UNIQUE (uc, tgl_ujian, mengulang_ke, mu)
);
CREATE INDEX ix_nilai_uc ON nilai(uc);

CREATE TABLE skl (
    id            INTEGER PRIMARY KEY,
    uc            TEXT NOT NULL,
    uc_raw        TEXT NOT NULL,
    noskl         TEXT,       -- NOT globally unique: collides across source files
    tgl_sidang    TEXT,
    tgl_cetak     TEXT,
    valid_through TEXT,
    nama          TEXT,
    diklat        TEXT,
    ijazah        TEXT,
    src           TEXT NOT NULL,
    UNIQUE (src, noskl)
);
CREATE INDEX ix_skl_uc ON skl(uc);

CREATE TABLE ref_ijzh (
    ijzh     TEXT PRIMARY KEY,
    kode     TEXT,
    deskripsi TEXT,
    mu_names TEXT NOT NULL     -- JSON list of active competency titles
);
CREATE TABLE ref_diklat (
    diklat    TEXT PRIMARY KEY,
    kode      TEXT,
    deskripsi TEXT
);

-- Feature 4: usage / abuse tracking. Never dropped.
CREATE TABLE IF NOT EXISTS access_log (
    id       INTEGER PRIMARY KEY,
    ts       TEXT NOT NULL,
    ip       TEXT,
    ua       TEXT,
    action   TEXT NOT NULL,   -- search | verify | view
    q        TEXT,            -- what was typed
    sc       TEXT,
    uc       TEXT,
    outcome  TEXT NOT NULL    -- ok | not_found | bad_verify | rate_limited
);
CREATE INDEX IF NOT EXISTS ix_log_ts ON access_log(ts);
CREATE INDEX IF NOT EXISTS ix_log_ip ON access_log(ip, ts);
"""

BAD = {"", "#N/A", "#REF!", "#VALUE!", "#NAME?", "0"}


def s(v):
    """Cell -> clean string, Excel error values become ''."""
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    t = str(v).strip()
    return "" if t in BAD else t


def d(v):
    """Cell -> ISO date string or None. Source stores 84% of dates as text."""
    if isinstance(v, datetime.datetime):
        v = v.date()
    if isinstance(v, datetime.date):
        return None if v.year < 1910 else v.isoformat()
    t = s(v)
    if not t:
        return None
    for f in ("%d %b %Y", "%d %B %Y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            p = datetime.datetime.strptime(t, f).date()
            return None if p.year < 1910 else p.isoformat()
        except ValueError:
            pass
    return None


def num(v):
    """Grade cell -> int or None. 2021 file holds 'U40', '40.', 'S 75', '--'."""
    if isinstance(v, (int, float)):
        return int(v)
    t = s(v)
    m = re.search(r"\d+", t)
    return int(m.group()) if m else None


def split_uc(uc):
    """'6212138564-03031N3120071-100' -> ('6212138564', '03031N3120071', '100').

    The sheets are inconsistent about the case of the exam number and about the
    trailing batch suffix, so both are normalised for matching.
    """
    parts = uc.split("-")
    if len(parts) < 2:
        return None
    sc = parts[0].strip()
    base = parts[1].strip().upper()
    suffix = parts[2].strip() if len(parts) > 2 else ""
    return sc, base, suffix


def canon(uc):
    """Canonical form of a UC key: exam number uppercased, suffix preserved."""
    p = split_uc(uc)
    if not p:
        return uc.upper()
    sc, base, suffix = p
    return f"{sc}-{base}-{suffix}" if suffix else f"{sc}-{base}"


def rows(ws, min_row):
    for r in ws.iter_rows(min_row=min_row, values_only=True):
        if any(v is not None for v in r):
            yield r


def load_cons(ws):
    """CONS holds 3 side-by-side lookup tables. Header row 3, data from row 4."""
    ijzh, diklat = {}, {}
    for r in rows(ws, 4):
        key = s(r[12]).upper()
        if key:
            names = [s(x) for x in r[15:37]]
            active = [n for n in names if n and n != "-"]
            ijzh[key] = (s(r[13]), s(r[14]), active)
        dk = s(r[4]).upper()
        if dk:
            diklat[dk] = (s(r[5]), s(r[6]))
    return ijzh, diklat


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=os.path.join(HERE, "ukp.db"))
    ap.add_argument("--src-dir", default=SRC_DIR)
    ap.add_argument("--exceptions", default=os.path.join(HERE, "exceptions.csv"))
    args = ap.parse_args()

    books = []
    for src, fname in SOURCES:
        path = os.path.join(args.src_dir, fname)
        if not os.path.exists(path):
            sys.exit(f"missing source: {path}")
        books.append((src, openpyxl.load_workbook(path, read_only=True, data_only=True)))

    con = sqlite3.connect(args.out)
    con.executescript(SCHEMA)
    stats, ijzh_ref = {}, {}

    # --- pass 1: reference tables + every person ---------------------------
    for src, wb in books:
        ij, dk = load_cons(wb["CONS"])
        ijzh_ref.update(ij)
        con.executemany("INSERT OR REPLACE INTO ref_ijzh VALUES (?,?,?,?)",
                        [(k, v[0], v[1], json.dumps(v[2])) for k, v in ij.items()])
        con.executemany("INSERT OR REPLACE INTO ref_diklat VALUES (?,?,?)",
                        [(k, v[0], v[1]) for k, v in dk.items()])

        n = 0
        for r in rows(wb["DataPeserta"], 2):
            uc, sc = canon(s(r[0])), s(r[1])
            if not uc or not sc or "-" not in uc:
                continue
            dob = d(r[5])
            if dob is None:  # fall back to the "CITY, 05 June 1996" text column
                m = re.search(r",\s*(\d{1,2}\s+[A-Za-z]+\s+\d{4})\s*$", s(r[9]))
                if m:
                    dob = d(m.group(1))
            usable = int(dob is not None
                         and datetime.date.fromisoformat(dob) not in PLACEHOLDER_DOB)
            con.execute("INSERT OR REPLACE INTO peserta VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                        (uc, sc, s(r[2]).upper(), s(r[3]).upper(), s(r[4]).upper(), dob,
                         usable, s(r[6]).upper(), s(r[7]).upper(), d(r[8]), src))
            n += 1
        stats.setdefault(src, {})["peserta"] = n
    con.commit()

    # Index for recovering keys whose batch suffix disagrees between sheets.
    # Only used when it points at exactly one person.
    by_base = {}
    for uc, sc in con.execute("SELECT uc, sc FROM peserta"):
        p = split_uc(uc)
        if p:
            by_base.setdefault((p[0], p[1]), set()).add(uc)

    known = {r[0] for r in con.execute("SELECT uc FROM peserta")}
    name_by_sc = {}
    for sc, nama in con.execute(
            "SELECT sc, nama FROM peserta WHERE nama<>'' GROUP BY sc"):
        name_by_sc[sc] = nama
    exceptions = []

    def resolve(raw, sheet, src, note):
        """Attach a record to a person, or log why it could not be attached."""
        uc = canon(raw)
        if uc in known:
            return uc
        p = split_uc(uc)
        if p:
            hit = by_base.get((p[0], p[1]))
            if hit and len(hit) == 1:
                return next(iter(hit))          # same exam, different batch suffix
            # the sheet's own NAMA is blank on most of these, so fall back to
            # the name registered against the seafarer code
            who = name_by_sc.get(p[0], "")
            if p[0] in {k[0] for k in by_base}:
                exceptions.append((src, sheet, raw, who,
                                   "SC dikenal, nomor ujian tidak cocok", note))
                return uc
            exceptions.append((src, sheet, raw, who,
                               "SC tidak ada di DataPeserta", note))
            return uc
        exceptions.append((src, sheet, raw, "", "kunci tidak dapat dibaca", note))
        return uc

    # --- pass 2: results and certificates ----------------------------------
    for src, wb in books:
        n = 0
        for r in rows(wb["DNILAI"], 2):
            raw = s(r[1])
            if not raw or "-" not in raw:
                continue
            ij = s(r[28]).upper()
            uc = resolve(raw, "DNILAI", src, f"{s(r[29]).upper()} {ij} {d(r[25]) or ''}".strip())
            width = len(ijzh_ref.get(ij, ("", "", []))[2]) or 22
            mu = [num(v) for v in r[2:24]][:width]
            got = [v for v in mu if v is not None]
            her_calc = sum(1 for v in got if v < 70)
            con.execute(
                "INSERT OR IGNORE INTO nilai (uc,uc_raw,ijzh,nama,tgl_ujian,tgl_sidang,"
                "mengulang_ke,her,her_calc,lulus,mu,src) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
                (uc, raw, ij, s(r[29]).upper(), d(r[25]), d(r[26]), num(r[27]), num(r[24]),
                 her_calc, int(bool(got) and her_calc == 0), json.dumps(mu), src))
            n += 1
        stats[src]["nilai"] = n

        # DataSKL: header is on row 5, rows 1-4 are a signature template
        n = 0
        for r in rows(wb["DataSKL"], 6):
            raw, noskl = s(r[1]), s(r[0])
            if not raw or "-" not in raw:
                continue  # drops the #REF! spill rows
            uc = resolve(raw, "DataSKL", src, f"SKL {noskl} {s(r[5]).upper()}".strip())
            con.execute(
                "INSERT OR IGNORE INTO skl (uc,uc_raw,noskl,tgl_sidang,tgl_cetak,"
                "valid_through,nama,diklat,ijazah,src) VALUES (?,?,?,?,?,?,?,?,?,?)",
                (uc, raw, noskl, d(r[2]), d(r[3]), d(r[4]), s(r[5]).upper(),
                 s(r[6]).upper(), s(r[7]).upper(), src))
            n += 1
        stats[src]["skl"] = n
        wb.close()
    con.commit()

    with open(args.exceptions, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["sumber", "sheet", "kunci_di_sheet", "nama_menurut_sc",
                    "masalah", "keterangan"])
        w.writerows(sorted(set(exceptions)))

    q = lambda sql: con.execute(sql).fetchone()[0]
    print("read per source:", json.dumps(stats))
    print("peserta   :", q("SELECT COUNT(*) FROM peserta"),
          " usable DOB:", q("SELECT COUNT(*) FROM peserta WHERE dob_usable=1"),
          " distinct SC:", q("SELECT COUNT(DISTINCT sc) FROM peserta"))
    print("nilai     :", q("SELECT COUNT(*) FROM nilai"),
          " lulus:", q("SELECT COUNT(*) FROM nilai WHERE lulus=1"),
          " her mismatch:", q("SELECT COUNT(*) FROM nilai WHERE her IS NOT NULL AND her<>her_calc"))
    print("skl       :", q("SELECT COUNT(*) FROM skl"),
          " distinct uc:", q("SELECT COUNT(DISTINCT uc) FROM skl"))
    print("ref_ijzh  :", q("SELECT COUNT(*) FROM ref_ijzh"),
          " ref_diklat:", q("SELECT COUNT(*) FROM ref_diklat"))
    print("recovered :", q("SELECT COUNT(*) FROM nilai WHERE uc<>uc_raw"), "nilai +",
          q("SELECT COUNT(*) FROM skl WHERE uc<>uc_raw"), "skl re-attached by key repair")
    print("latest    : ujian", q("SELECT MAX(tgl_ujian) FROM nilai"),
          " sidang", q("SELECT MAX(tgl_sidang) FROM nilai"),
          " cetak", q("SELECT MAX(tgl_cetak) FROM skl"))
    print("exceptions:", len(set(exceptions)), "->", args.exceptions)

    # self-check: the invariants the app relies on
    orphan = q("SELECT COUNT(*) FROM nilai n LEFT JOIN peserta p USING(uc) WHERE p.uc IS NULL")
    assert q("SELECT COUNT(*) FROM peserta") > 30000, "peserta under-loaded"
    assert q("SELECT COUNT(*) FROM nilai") > 39000, "nilai under-loaded (union failed?)"
    assert q("SELECT COUNT(*) FROM skl") > 7000, "skl under-loaded"
    assert q("SELECT COUNT(*) FROM ref_ijzh") > 40, "CONS lookup not parsed"
    assert orphan < 30, f"too many nilai rows with no peserta: {orphan}"
    unattached = {r[0] for r in con.execute(
        "SELECT uc_raw FROM nilai n LEFT JOIN peserta p USING(uc) WHERE p.uc IS NULL"
        " UNION SELECT uc_raw FROM skl s LEFT JOIN peserta p USING(uc) WHERE p.uc IS NULL")}
    reported = {e[2] for e in exceptions}
    assert unattached == reported, (
        f"exception report out of sync: {unattached ^ reported}")
    print("self-check OK (orphan nilai rows: %d, unattached keys: %d)"
          % (orphan, len(unattached)))
    con.close()


if __name__ == "__main__":
    main()
